from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditLog, Unit
from tests.conftest import sign_in

CS = "/api/teaching/CS"
ED = "/api/teaching/ED"


def _ok(response):
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _course(client, base, programme, students, year=1, semester=1, **extra):
    body = {"programme": programme, "year": year, "semester": semester, "students": students, **extra}
    return _ok(client.post(f"{base}/courses", json=body))["id"]


def _lecturer(client, base, name, staff, **extra):
    body = {"full_name": name, "title": "Dr.", "staff_number": staff, **extra}
    return _ok(client.post(f"{base}/lecturers", json=body))["id"]


def _unit(client, base, code, course_ids, hours=3, length=3, **extra):
    body = {
        "code": code,
        "title": f"{code} title",
        "weekly_hours": hours,
        "class_length": length,
        "course_ids": course_ids,
        **extra,
    }
    return _ok(client.post(f"{base}/units", json=body))["id"]


def _board(client, base=CS):
    return _ok(client.get(base))


def _whole_group(client, base, unit_id):
    unit = next(u for u in _board(client, base)["units"] if u["id"] == unit_id)
    return unit["groups"][0]["id"]


def _ready_cs_unit(client):
    """CS enters a course, a lecturer and a whole-class unit taught by them: ready to publish."""
    course = _course(client, CS, "BSC.CS", 120)
    jane = _lecturer(client, CS, "Jane Otieno", "stf/9", max_weekly_hours=18)
    unit = _unit(client, CS, "cosc 103", [course])
    _ok(
        client.patch(
            f"{CS}/groups/{_whole_group(client, CS, unit)}", json={"lecturer_id": jane, "lecture_hours": 30}
        )
    )
    return course, jane, unit


def test_a_department_enters_courses_units_and_groups(client, generator):
    sign_in(client, "STF/0101")  # CS allocation staff
    _ready_cs_unit(client)
    board = _board(client)
    assert [c["label"] for c in board["courses"]] == ["BSC.CS Y1S1"]
    [row] = board["units"]
    assert row["code"] == "COSC 103" and row["servicing"] is False
    [group] = row["groups"]
    assert group["label"] == "" and group["key"] == "COSC103-MAIN"
    assert group["students"] == 120 and group["lecturer"] == "Dr. Jane Otieno (FT)"
    assert board["lecturers"][0]["weekly_hours"] == 3
    assert board["problems"] == []
    assert board["can_edit"] is True and board["can_publish"] is False  # staff can't publish


def test_a_split_must_add_up_to_the_course(client, generator):
    sign_in(client, "STF/0100")
    course = _course(client, CS, "BSC.CS", 120)
    unit = _unit(client, CS, "COSC 103", [course])
    split = [
        {"label": "A", "members": [{"course_id": course, "students": 60}]},
        {"label": "B", "members": [{"course_id": course, "students": 50}]},
    ]
    _ok(client.put(f"{CS}/units/{unit}/groups", json={"groups": split}))
    board = _board(client)
    assert [g["label"] for g in board["units"][0]["groups"]] == ["GR A", "GR B"]
    assert any(p["code"] == "split_total" and "120 students" in p["message"] for p in board["problems"])
    split[1]["members"][0]["students"] = 60
    _ok(client.put(f"{CS}/units/{unit}/groups", json={"groups": split}))
    assert not any(p["code"] == "split_total" for p in _board(client)["problems"])
    # A split names every group.
    bad = client.put(f"{CS}/units/{unit}/groups", json={"groups": [split[0], {**split[1], "label": ""}]})
    assert bad.status_code == 422


def test_servicing_picks_another_departments_course(client, generator):
    sign_in(client, "STF/0103")  # Education's COD
    edu_course = _course(client, ED, "BED.SCI", 80)
    client.cookies.clear()
    sign_in(client, "STF/0100")  # CS's COD
    board = _board(client)
    assert [c["label"] for c in board["other_courses"]] == ["BED.SCI Y1S1"]
    unit = _unit(client, CS, "COSC 110", [edu_course])
    assert _board(client)["units"][0]["servicing"] is True
    # CS can read Education's courses, but not change them.
    assert client.patch(f"{CS}/courses/{edu_course}", json={"students": 10}).status_code == 403
    # Education can't enter a unit CS already entered, or drop a course CS's unit takes.
    client.cookies.clear()
    sign_in(client, "STF/0103")
    taken = client.post(
        f"{ED}/units",
        json={
            "code": "COSC110",
            "title": "Intro",
            "weekly_hours": 3,
            "class_length": 3,
            "course_ids": [edu_course],
        },
    )
    assert taken.status_code == 409 and "already entered by CS" in taken.json()["error"]["message"]
    in_use = client.delete(f"{ED}/courses/{edu_course}")
    assert in_use.status_code == 409 and "COSC 110 (CS)" in in_use.json()["error"]["message"]
    assert unit


def test_a_cohort_is_entered_once(client, generator):
    sign_in(client, "STF/0100")
    _course(client, CS, "BSC.CS", 120)
    again = client.post(
        f"{CS}/courses", json={"programme": "bsc.cs", "year": 1, "semester": 1, "students": 5}
    )
    assert again.status_code == 409 and again.json()["error"]["code"] == "COURSE_TAKEN"


def test_publish_locks_until_the_cod_reopens(client, generator, db):
    sign_in(client, "STF/0100")
    _course, _jane, unit = _ready_cs_unit(client)
    board = _ok(client.post(f"{CS}/publish"))
    assert board["published"] is True and board["can_edit"] is False
    assert board["publication"]["unit_count"] == 1
    locked = client.patch(f"{CS}/units/{unit}", json={"title": "New title"})
    assert locked.status_code == 409 and locked.json()["error"]["code"] == "PUBLISHED"
    assert client.post(f"{CS}/publish").status_code == 409

    board = _ok(client.post(f"{CS}/reopen"))
    assert board["published"] is False and board["publication"]["reopened_at"]
    _ok(client.patch(f"{CS}/units/{unit}", json={"weekly_hours": 4}))
    board = _board(client)
    assert board["units"][0]["changed_since_published"] is True and board["summary"]["changed"] == 1
    board = _ok(client.post(f"{CS}/publish"))
    assert board["units"][0]["changed_since_published"] is False
    actions = [a.action for a in db.scalars(select(AuditLog).order_by(AuditLog.at))]
    assert actions.count("department.publish") == 2 and "department.reopen" in actions


def test_only_the_cod_publishes_and_nothing_can_be_missing(client, generator):
    sign_in(client, "STF/0101")
    course = _course(client, CS, "BSC.CS", 120)
    _unit(client, CS, "COSC 103", [course])
    assert client.post(f"{CS}/publish").status_code == 403
    client.cookies.clear()
    sign_in(client, "STF/0100")
    refused = client.post(f"{CS}/publish")
    assert refused.status_code == 409
    codes = {p["code"] for p in refused.json()["error"]["details"]["problems"]}
    assert codes == {"no_lecturer"}


def test_a_lecturer_needs_a_staff_number_to_be_published(client, generator, db):
    sign_in(client, "STF/0100")
    _course, jane, _unit = _ready_cs_unit(client)
    _ok(client.patch(f"{CS}/lecturers/{jane}", json={"staff_number": None}))
    assert any(p["code"] == "no_staff_number" for p in _board(client)["problems"])


def test_workload_counts_every_department_for_one_person(client, generator):
    sign_in(client, "STF/0103")
    edu_course = _course(client, ED, "BED.SCI", 80)
    shared = _lecturer(client, ED, "Sam Mwangi", "STF/77", max_weekly_hours=6)
    edu_unit = _unit(client, ED, "EDFO 111", [edu_course], hours=4, length=2)
    _ok(client.patch(f"{ED}/groups/{_whole_group(client, ED, edu_unit)}", json={"lecturer_id": shared}))
    assert _board(client, ED)["problems"] == []
    client.cookies.clear()
    sign_in(client, "STF/0100")
    course = _course(client, CS, "BSC.CS", 120)
    same_person = _lecturer(client, CS, "Sam Mwangi", "stf/77", max_weekly_hours=6)
    unit = _unit(client, CS, "COSC 103", [course], hours=4, length=2)
    _ok(client.patch(f"{CS}/groups/{_whole_group(client, CS, unit)}", json={"lecturer_id": same_person}))
    board = _board(client)
    assert board["lecturers"][0]["weekly_hours"] == 8
    [problem] = board["problems"]
    assert problem["code"] == "over_hours" and "8 hours a week across all departments" in problem["message"]


def test_a_group_takes_only_its_own_departments_active_lecturers(client, generator):
    sign_in(client, "STF/0103")
    other = _lecturer(client, ED, "Eve Other", "STF/50")
    client.cookies.clear()
    sign_in(client, "STF/0100")
    course = _course(client, CS, "BSC.CS", 120)
    unit = _unit(client, CS, "COSC 103", [course])
    group = _whole_group(client, CS, unit)
    assert client.patch(f"{CS}/groups/{group}", json={"lecturer_id": other}).status_code == 403
    jane = _lecturer(client, CS, "Jane Otieno", "STF/9")
    _ok(client.patch(f"{CS}/groups/{group}", json={"lecturer_id": jane}))
    in_use = client.delete(f"{CS}/lecturers/{jane}")
    assert in_use.status_code == 409 and in_use.json()["error"]["code"] == "IN_USE"


def test_other_departments_are_closed(client, generator):
    sign_in(client, "STF/0100")
    assert client.get(ED).status_code == 403
    assert (
        client.post(
            f"{ED}/courses", json={"programme": "X", "year": 1, "semester": 1, "students": 1}
        ).status_code
        == 403
    )


def test_the_directorate_sees_every_department_read_only(client, generator):
    sign_in(client, "STF/0100")
    _ready_cs_unit(client)
    client.post(f"{CS}/publish")
    client.cookies.clear()
    sign_in(client, "STF/0102")
    rows = {d["code"]: d for d in _ok(client.get("/api/teaching/progress"))["departments"]}
    assert rows["CS"]["published"] is True and rows["CS"]["units"] == 1
    assert rows["ED"]["published"] is False
    board = _board(client)
    assert board["role"] is None and board["can_edit"] is False
    assert (
        client.post(
            f"{CS}/courses", json={"programme": "X", "year": 1, "semester": 1, "students": 1}
        ).status_code
        == 403
    )


# -- what the Generator reads ---------------------------------------------------------------


def test_submissions_need_the_generators_key(client, generator):
    url = "/api/integrations/generator/terms/term-1/submissions"
    assert client.get(url).status_code == 401
    assert client.get(url, headers={"X-API-Key": "wrong"}).status_code == 401
    assert client.get(url, headers={"X-API-Key": "test-inbound-key"}).status_code == 200


def test_submissions_carry_published_departments_and_the_courses_they_take(client, generator):
    sign_in(client, "STF/0103")
    edu_course = _course(client, ED, "BED.SCI", 80)
    client.cookies.clear()
    sign_in(client, "STF/0100")
    course = _course(client, CS, "BSC.CS", 120)
    jane = _lecturer(client, CS, "Jane Otieno", "stf/9")
    unit = _unit(client, CS, "COSC 103", [course, edu_course], hours=3, length=2)
    split = [
        {"label": "A", "members": [{"course_id": course}], "lecturer_id": jane, "practical_hours": 15},
        {"label": "B", "members": [{"course_id": edu_course}], "lecturer_id": jane},
    ]
    _ok(client.put(f"{CS}/units/{unit}/groups", json={"groups": split}))
    _ok(client.post(f"{CS}/publish"))

    url = "/api/integrations/generator/terms/term-1/submissions"
    data = _ok(client.get(url, headers={"X-API-Key": "test-inbound-key"}))
    assert sorted(c["label"] for c in data["courses"]) == ["BED.SCI Y1S1", "BSC.CS Y1S1"]
    [cs] = data["published"]
    assert cs["code"] == "CS" and cs["published_at"]
    [unit_out] = cs["units"]
    assert unit_out["code"] == "COSC 103" and unit_out["class_length"] == 2
    group_a = unit_out["groups"][0]
    assert group_a["key"] == "COSC103-GRA" and group_a["students"] == 120
    assert group_a["lecturer"] == {
        "staff_number": "STF/9",
        "name": "Dr. Jane Otieno (FT)",
        "employment_type": "full_time",
    }
    assert group_a["practical_hours"] == 15
    assert {d["code"] for d in data["not_published"]} == {"ED"}

    # Reopened, CS is waiting again, with what changed.
    _ok(client.post(f"{CS}/reopen"))
    _ok(client.patch(f"{CS}/units/{unit}", json={"title": "Programming"}))
    data = _ok(client.get(url, headers={"X-API-Key": "test-inbound-key"}))
    assert data["published"] == []
    cs_waiting = next(d for d in data["not_published"] if d["code"] == "CS")
    assert cs_waiting["changed_units"] == ["COSC 103"] and cs_waiting["reopened_at"]


# -- times, once the Director publishes ---------------------------------------------------------


def test_published_times_show_on_the_page_and_downloads(client, generator):
    sign_in(client, "STF/0100")
    _ready_cs_unit(client)
    assert _board(client)["units"][0]["groups"][0]["meetings"] == []
    generator.publish({"COSC103-MAIN": [("TUE", "10:00", "13:00", "STB 1")]})
    board = _board(client)
    assert board["times"]["published_version"] == 1
    assert board["units"][0]["groups"][0]["meetings"][0]["room"] == "STB 1"
    response = client.get(f"{CS}/export.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    cells = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
    assert "Day, time & room" in cells and "TUE 10:00-13:00 STB 1" in cells
    for fmt in ("pdf", "docx"):
        assert client.get(f"{CS}/export.{fmt}").status_code == 200


# -- paste from Excel -------------------------------------------------------------------------


def test_the_workbook_tabs_fill_courses_units_and_groups(client, generator, db):
    sign_in(client, "STF/0103")
    _course(client, ED, "BED.SCI", 80)
    client.cookies.clear()
    sign_in(client, "STF/0100")
    result = _ok(
        client.post(
            f"{CS}/courses/bulk",
            json={
                "rows": [
                    {"Programme": "BSC.CS", "Year": "1", "Semester": "1", "Students": "120"},
                    {"Course": "BSC.CS Y2S1 (A)", "Students": "40"},
                    {"Programme": "BSC.CS", "Year": "9", "Semester": "1", "Students": "1"},
                ]
            },
        )
    )
    assert result["created"] == 2 and result["errors"][0]["line"] == 3
    result = _ok(
        client.post(
            f"{CS}/units/bulk",
            json={
                "rows": [
                    {
                        "Code": "COSC 103",
                        "Title": "Programming",
                        "Hours per week": "3",
                        "Class length": "2",
                        "Room type": "Lab",
                        "Courses": "BSC.CS Y1S1; BED.SCI Y1S1",
                    },
                    {"Code": "COSC 200", "Title": "Data", "Hours per week": "3", "Courses": "NOPE Y1S1"},
                ]
            },
        )
    )
    assert result["created"] == 1 and "NOPE Y1S1" in result["errors"][0]["message"]
    result = _ok(
        client.post(
            f"{CS}/groups/bulk",
            json={
                "rows": [
                    {
                        "Unit code": "COSC103",
                        "Group": "A",
                        "Courses": "BSC.CS Y1S1",
                        "Students": "60",
                        "Lecturer": "Dr. Jane Otieno (PT)",
                        "Staff no.": "STF/9",
                        "L": "30",
                    },
                    {
                        "Unit code": "COSC103",
                        "Group": "B",
                        "Courses": "BSC.CS Y1S1; BED.SCI Y1S1",
                        "Students": "60; 80",
                        "Staff no.": "STF/9",
                    },
                ]
            },
        )
    )
    assert result["errors"] == [] and result["groups"] == 2 and result["lecturers_added"] == 1
    unit = db.scalar(select(Unit).where(Unit.code_key == "COSC103"))
    assert unit is not None and unit.room_type == "lab"
    board = _board(client)
    groups = board["units"][0]["groups"]
    assert [(g["label"], g["students"]) for g in groups] == [("GR A", 60), ("GR B", 140)]
    assert groups[0]["lecturer"] == "Dr. Jane Otieno (PT)" and groups[0]["lecture_hours"] == 30
    assert board["problems"] == []
    # Pasting the units again updates in place.
    again = _ok(
        client.post(
            f"{CS}/units/bulk",
            json={
                "rows": [
                    {"Code": "cosc 103", "Title": "Programming I", "Hours per week": "4", "Class length": "2"}
                ]
            },
        )
    )
    assert again["updated"] == 1 and _board(client)["units"][0]["title"] == "Programming I"


def test_a_bad_groups_row_leaves_the_unit_as_it_was(client, generator):
    sign_in(client, "STF/0100")
    course = _course(client, CS, "BSC.CS", 120)
    _unit(client, CS, "COSC 103", [course])
    result = _ok(
        client.post(
            f"{CS}/groups/bulk",
            json={
                "rows": [
                    {"Unit code": "COSC 103", "Group": "A", "Students": "60"},
                    {"Unit code": "COSC 103", "Group": "B", "Students": "9999"},
                ]
            },
        )
    )
    assert result["groups"] == 0 and result["error_count"] == 1
    assert [g["label"] for g in _board(client)["units"][0]["groups"]] == [""]
