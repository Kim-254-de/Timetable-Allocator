from __future__ import annotations

from io import BytesIO

from docx import Document

from tests.conftest import sign_in

CS = "/api/teaching/CS"
ED = "/api/teaching/ED"
KEY = {"X-API-Key": "test-inbound-key"}


def _ok(response):
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _course(client, base, programme, students, **extra):
    body = {"programme": programme, "year": 1, "semester": 1, "students": students, **extra}
    return _ok(client.post(f"{base}/courses", json=body))["id"]


def _unit(client, base, code, courses, lecturer=None):
    body = {
        "code": code,
        "title": f"{code} title",
        "weekly_hours": 3,
        "class_length": 3,
        "course_ids": courses,
    }
    unit = _ok(client.post(f"{base}/units", json=body))["id"]
    if lecturer:
        board = _ok(client.get(base))
        group = next(u for u in board["units"] if u["id"] == unit)["groups"][0]["id"]
        _ok(client.patch(f"{base}/groups/{group}", json={"lecturer_id": lecturer}))
    return unit


def _lecturer(client, base, name, staff):
    return _ok(
        client.post(f"{base}/lecturers", json={"full_name": name, "title": "Dr.", "staff_number": staff})
    )["id"]


def _exams(client, base=CS):
    return _ok(client.get(f"{base}/exams"))


def _setup(client):
    """Education teaches BITE 131 to its course; CS teaches COSC 130 to its own and an ODEL course."""
    sign_in(client, "STF/0103")
    ed_course = _course(client, ED, "BED.SCI", 80)
    ed_unit = _unit(client, ED, "BITE 131", [ed_course])
    client.cookies.clear()
    sign_in(client, "STF/0100")
    cs_course = _course(client, CS, "BSC.CS", 120)
    odel = _course(client, CS, "BSC.CS", 30, sub_class="ODEL", mode="odel")
    jane = _lecturer(client, CS, "Jane Otieno", "STF/9")
    cosc130 = _unit(client, CS, "COSC 130", [cs_course, odel], lecturer=jane)
    cosc200 = _unit(client, CS, "COSC 200", [cs_course], lecturer=jane)
    return {"ed_unit": ed_unit, "cosc130": cosc130, "cosc200": cosc200, "jane": jane}


def test_a_shared_paper_is_one_paper(client, generator):
    ids = _setup(client)
    _ok(
        client.post(
            f"{CS}/exam-papers", json={"unit_ids": [ids["cosc130"], ids["ed_unit"]], "title": "Programming"}
        )
    )
    [paper] = _exams(client)["papers"]
    assert paper["codes"] == "BITE 131/COSC 130"
    assert paper["examiner"] == "Dr. Jane Otieno (FT)"  # the unit's lecturer
    assert paper["students"] == 230 and sorted(paper["class_rows"]) == [
        "BED.SCI Y1S1",
        "BSC.CS Y1S1",
        "BSC.CS Y1S1 (ODEL)",
    ]
    # Education sees its unit is on CS's paper, and can't put it on another.
    client.cookies.clear()
    sign_in(client, "STF/0103")
    [unit] = _exams(client, ED)["units"]
    assert unit["paper_codes"] == "BITE 131/COSC 130" and unit["paper_department"] == "CS"
    taken = client.post(f"{ED}/exam-papers", json={"unit_ids": [ids["ed_unit"]]})
    assert taken.status_code == 409 and taken.json()["error"]["code"] == "UNIT_ON_PAPER"


def test_a_paper_needs_one_of_the_departments_units(client, generator):
    ids = _setup(client)
    refused = client.post(f"{CS}/exam-papers", json={"unit_ids": [ids["ed_unit"]]})
    assert refused.status_code == 422


def test_duplicate_codes_are_refused_when_pasted(client, generator):
    _setup(client)
    result = _ok(
        client.post(
            f"{CS}/exam-papers/bulk",
            json={
                "rows": [
                    {"Codes": "COSC 130", "Title": "Programming"},
                    {"Codes": "COSC 200/COSC 130", "Title": "Again"},
                    {"Codes": "COSC 999"},
                ]
            },
        )
    )
    assert result["created"] == 1 and result["error_count"] == 2
    assert "also on row 1" in result["errors"][0]["message"] and "COSC 999" in result["errors"][1]["message"]


def test_publishing_needs_every_unit_examined(client, generator):
    ids = _setup(client)
    assert _ok(client.post(f"{CS}/exam-papers/from-units"))["created"] == 2
    board = _exams(client)
    assert board["problems"] == []
    paper200 = next(p for p in board["papers"] if p["codes"] == "COSC 200")
    _ok(client.delete(f"{CS}/exam-papers/{paper200['id']}"))
    refused = client.post(f"{CS}/exams/publish")
    assert refused.status_code == 409 and "COSC 200 is on no paper" in refused.json()["error"]["message"]
    _ok(client.put(f"{CS}/units/{ids['cosc200']}/no-written-exam", json={"no_written_exam": True}))
    board = _ok(client.post(f"{CS}/exams/publish"))
    assert board["published"] is True and board["can_edit"] is False
    assert client.post(f"{CS}/exam-papers/from-units").status_code == 409  # locked
    # The exam list and the teaching list are published separately.
    assert _ok(client.get(CS))["published"] is False
    board = _ok(client.post(f"{CS}/exams/reopen"))
    assert board["published"] is False


def test_practicals_are_on_weekdays(client, generator):
    ids = _setup(client)
    body = {"unit_id": ids["cosc200"], "date": "2026-11-21", "start_time": "09:00", "end_time": "11:00"}
    assert client.post(f"{CS}/practicals", json=body).status_code == 422  # a Saturday
    _ok(client.post(f"{CS}/practicals", json={**body, "date": "2026-11-20"}))
    pasted = _ok(
        client.post(
            f"{CS}/practicals/bulk",
            json={
                "rows": [{"Unit code": "COSC 130", "Date": "19/11/2026", "Start": "2:00pm", "End": "16:00"}]
            },
        )
    )
    assert pasted["created"] == 1
    practicals = _exams(client)["practicals"]
    assert [(p["unit_code"], p["date"], p["start_time"]) for p in practicals] == [
        ("COSC 130", "2026-11-19", "14:00"),
        ("COSC 200", "2026-11-20", "09:00"),
    ]


def test_the_moderation_letter_lists_the_papers(client, generator):
    ids = _setup(client)
    _ok(
        client.post(
            f"{CS}/exam-papers", json={"unit_ids": [ids["cosc130"], ids["ed_unit"]], "title": "Programming"}
        )
    )
    response = client.get(
        f"{CS}/moderation-letter.docx",
        params={"moderator": "Prof. A Moderator\nP.O. Box 1, Nairobi", "date": "1 Nov 2026"},
    )
    assert response.status_code == 200
    doc = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "RE: MODERATION OF SEPTEMBER - DECEMBER 2026 COMPUTER SCIENCE EXAMINATION DRAFTS" in text
    assert "Prof. A Moderator" in text and "CC: Registrar A&A" in text
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert cells[:6] == ["S/NO", "UNIT CODE(S)", "UNIT TITLE", "1", "BITE 131/COSC 130", "Programming"]
    assert client.get(f"{CS}/moderation-letter.pdf").content[:4] == b"%PDF"


def test_the_generator_reads_published_papers_and_practicals(client, generator):
    ids = _setup(client)
    _ok(
        client.post(
            f"{CS}/exam-papers", json={"unit_ids": [ids["cosc130"], ids["ed_unit"]], "title": "Programming"}
        )
    )
    _ok(client.post(f"{CS}/exam-papers", json={"unit_ids": [ids["cosc200"]]}))
    body = {"unit_id": ids["cosc200"], "date": "2026-11-20", "start_time": "09:00", "end_time": "11:00"}
    _ok(client.post(f"{CS}/practicals", json=body))
    url = "/api/integrations/generator/terms/term-1/exams"
    assert client.get(url).status_code == 401
    assert _ok(client.get(url, headers=KEY))["published"] == []
    _ok(client.post(f"{CS}/exams/publish"))
    data = _ok(client.get(url, headers=KEY))
    [cs] = data["published"]
    shared = next(p for p in cs["papers"] if p["codes"] == ["BITE 131", "COSC 130"])
    assert shared["students"] == 230 and shared["examiner"]["staff_number"] == "STF/9"
    assert sorted(r["mode"] for r in shared["rows"]) == ["odel", "regular", "regular"]
    assert cs["practicals"][0]["codes"] == ["COSC 200"] and cs["practicals"][0]["date"] == "2026-11-20"
    assert {c["label"] for c in data["courses"]} == {"BED.SCI Y1S1", "BSC.CS Y1S1", "BSC.CS Y1S1 (ODEL)"}
    assert {d["code"] for d in data["not_published"]} == {"ED"}
