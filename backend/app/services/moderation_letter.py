"""
The letter a COD sends with the department's exam drafts for moderation, as
Word and PDF, built from the paper list (replacing the typed letter):

    THARAKA UNIVERSITY
    DEPARTMENT OF COMPUTER SCIENCE AND ICT
    <date>
    <moderator's name and address>
    RE: MODERATION OF <TERM> <DEPARTMENT> EXAMINATION DRAFTS
    <covering paragraph>
    S/NO | UNIT CODE(S) | UNIT TITLE
    ...
    <COD's signature block>
    CC: Registrar A&A
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

UNIVERSITY = "THARAKA UNIVERSITY"
HEADERS = ["S/NO", "UNIT CODE(S)", "UNIT TITLE"]


@dataclass
class Letter:
    department: str
    term: str
    date: str
    moderator: list[str]  # name, then address lines
    papers: list[tuple[str, str]]  # (codes, title)
    cod: str

    @property
    def heading(self) -> str:
        return f"RE: MODERATION OF {self.term.upper()} {self.department.upper()} EXAMINATION DRAFTS"

    @property
    def body(self) -> str:
        return (
            f"Please find enclosed {len(self.papers)} examination drafts from the {self.department} "
            f"for the {self.term} examinations, listed below, for your moderation."
        )

    def rows(self) -> list[list[str]]:
        return [[str(i), codes, title] for i, (codes, title) in enumerate(self.papers, start=1)]


def build_docx(letter: Letter) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.5)
    doc.styles["Normal"].font.size = Pt(11)
    for line in (UNIVERSITY, letter.department.upper()):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(13)
    doc.add_paragraph(letter.date)
    for line in letter.moderator:
        doc.add_paragraph(line).paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    doc.add_paragraph().add_run(letter.heading).bold = True
    doc.add_paragraph(letter.body)
    table = doc.add_table(rows=len(letter.papers) + 1, cols=3)
    table.style = "Table Grid"
    widths = [Cm(1.5), Cm(5), Cm(10)]
    for r, values in enumerate([HEADERS, *letter.rows()]):
        for c, value in enumerate(values):
            cell = table.rows[r].cells[c]
            cell.text = value
            cell.width = widths[c]
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()
    doc.add_paragraph("Yours faithfully,")
    doc.add_paragraph()
    doc.add_paragraph(letter.cod)
    doc.add_paragraph(f"Chair, {letter.department}")
    doc.add_paragraph()
    doc.add_paragraph("CC: Registrar A&A")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def build_pdf(letter: Letter) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    title = ParagraphStyle("title", parent=styles["Title"], fontSize=13, leading=16, spaceAfter=0)
    body = styles["BodyText"]
    cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=11)
    story: list[Any] = [
        Paragraph(escape(UNIVERSITY), title),
        Paragraph(escape(letter.department.upper()), title),
    ]
    story += [Spacer(1, 6 * mm), Paragraph(escape(letter.date), body), Spacer(1, 3 * mm)]
    story += [Paragraph(escape(line), body) for line in letter.moderator]
    story += [
        Spacer(1, 4 * mm),
        Paragraph(f"<b>{escape(letter.heading)}</b>", body),
        Paragraph(escape(letter.body), body),
    ]
    data = [HEADERS] + [[Paragraph(escape(v), cell) for v in row] for row in letter.rows()]
    table = Table(data, colWidths=[14 * mm, 45 * mm, 115 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story += [
        Spacer(1, 3 * mm),
        table,
        Spacer(1, 8 * mm),
        Paragraph("Yours faithfully,", body),
        Spacer(1, 10 * mm),
    ]
    story += [Paragraph(escape(letter.cod), body), Paragraph(escape(f"Chair, {letter.department}"), body)]
    story += [Spacer(1, 6 * mm), Paragraph("CC: Registrar A&amp;A", body)]
    out = BytesIO()
    SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=letter.heading,
    ).build(story)
    return out.getvalue()


BUILDERS = {"docx": build_docx, "pdf": build_pdf}
MEDIA = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
