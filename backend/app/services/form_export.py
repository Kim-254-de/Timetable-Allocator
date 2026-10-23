"""
The department's allocation form as PDF, Word and Excel, laid out like the
Word form departments send today:

    THARAKA UNIVERSITY
    DEPARTMENT OF COMPUTER SCIENCE AND ICT
    SEPTEMBER - DECEMBER 2026 COURSE ALLOCATION

    BSC.COMP SCI Y1S1 (120 Students)
    Course code | Course title | L | P | CF | Lecturer (FT/PT) | Phone [| Day, time & room]
    ...
    TOTAL                        L   P   CF

    SERVICING FOR OTHER DEPARTMENTS
    ...
    Prepared by: <name>, COD

Phone numbers are printed only here, on the department's own download. Once
the Director has published the timetable, each class's day, time and room is
printed too.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

UNIVERSITY = "THARAKA UNIVERSITY"
HEADERS = ["Course code", "Course title", "L", "P", "CF", "Lecturer (FT/PT)", "Phone"]
WHEN = "Day, time & room"


@dataclass
class Block:
    heading: str
    rows: list[list[str]] = field(default_factory=list)
    totals: list[float] = field(default_factory=lambda: [0.0, 0.0, 0.0])


@dataclass
class Layout:
    title_lines: list[str]
    own: list[Block]
    servicing: list[Block]
    prepared_by: str
    with_times: bool = False

    @property
    def headers(self) -> list[str]:
        return [*HEADERS, WHEN] if self.with_times else HEADERS


def _h(value: float | None) -> str:
    return "" if value is None else f"{value:g}"


def _when(meetings: list[dict[str, str]]) -> str:
    return "; ".join(
        f"{m['day_of_week']} {m['start_time']}-{m['end_time']}" + (f" {m['room']}" if m.get("room") else "")
        for m in meetings
    )


def layout(board: dict[str, Any], prepared_by: str) -> Layout:
    """One row per class (a unit, or one group of a split unit), under the first course that takes it."""
    term = board["term"]
    term_name = str(term.get("name") or term.get("label") or "").upper()
    title = [UNIVERSITY, str(board["department"]["name"]).upper(), f"{term_name} COURSE ALLOCATION".strip()]
    phones = {lec["id"]: lec.get("phone") or "" for lec in board.get("lecturers") or []}
    with_times = any(g["meetings"] for u in board["units"] for g in u["groups"])
    own: dict[str, Block] = {}
    servicing: dict[str, Block] = {}
    for unit in board["units"]:
        for group in unit["groups"]:
            members = group["members"] or [{"label": "Other units", "students": None}]
            first = members[0]
            heading = first["label"] + (f" ({first['students']} Students)" if first.get("students") else "")
            target = servicing if unit["servicing"] else own
            block = target.setdefault(heading, Block(heading))
            others = [m["label"] for m in members[1:]]
            unit_title = unit["title"] + (f" (with {', '.join(others)})" if others else "")
            hours = [group.get("lecture_hours"), group.get("practical_hours"), group.get("cf_hours")]
            row = [
                f"{unit['code']} {group['label']}".strip(),
                unit_title,
                *[_h(h) for h in hours],
                group.get("lecturer") or "To be allocated",
                phones.get(group.get("lecturer_id") or "", ""),
            ]
            if with_times:
                row.append(_when(group["meetings"]) or "Not timetabled")
            block.rows.append(row)
            for i, h in enumerate(hours):
                block.totals[i] += float(h or 0)
    return Layout(
        title,
        sorted(own.values(), key=lambda b: b.heading),
        sorted(servicing.values(), key=lambda b: b.heading),
        prepared_by,
        with_times,
    )


def _sections(doc: Layout) -> list[tuple[str | None, list[Block]]]:
    out: list[tuple[str | None, list[Block]]] = [(None, doc.own)]
    if doc.servicing:
        out.append(("SERVICING FOR OTHER DEPARTMENTS", doc.servicing))
    return out


def _total_row(doc: Layout, block: Block) -> list[str]:
    return ["TOTAL", "", *[_h(t) for t in block.totals], ""] + [""] * (len(doc.headers) - 6)


# -- PDF --------------------------------------------------------------------------------


def build_pdf(doc: Layout) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    title = ParagraphStyle("title", parent=styles["Title"], fontSize=12, leading=15, spaceAfter=0)
    heading = ParagraphStyle("heading", parent=styles["Heading4"], spaceBefore=8, spaceAfter=4)
    section = ParagraphStyle("section", parent=styles["Heading3"], spaceBefore=12)

    story: list[Any] = [Paragraph(escape(line), title) for line in doc.title_lines]
    story.append(Spacer(1, 6 * mm))
    widths = (
        [22 * mm, 40 * mm, 8 * mm, 8 * mm, 8 * mm, 40 * mm, 22 * mm, 36 * mm]
        if doc.with_times
        else [28 * mm, 54 * mm, 9 * mm, 9 * mm, 9 * mm, 51 * mm, 24 * mm]
    )
    for name, blocks in _sections(doc):
        if name:
            story.append(Paragraph(name, section))
        for block in blocks:
            story.append(Paragraph(escape(block.heading), heading))
            data = (
                [doc.headers]
                + [[Paragraph(escape(c), cell) for c in r] for r in block.rows]
                + [_total_row(doc, block)]
            )
            table = Table(data, colWidths=widths, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
    if not doc.own and not doc.servicing:
        story.append(Paragraph("No units entered yet.", cell))
    story.append(Spacer(1, 10 * mm))
    story.append(Paragraph(escape(f"Prepared by: {doc.prepared_by}, COD"), styles["BodyText"]))
    out = BytesIO()
    SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=doc.title_lines[-1],
    ).build(story)
    return out.getvalue()


# -- Word -------------------------------------------------------------------------------


def build_docx(doc: Layout) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    word = Document()
    for s in word.sections:
        s.left_margin = s.right_margin = Cm(1.5)
        s.top_margin = s.bottom_margin = Cm(1.5)
    word.styles["Normal"].font.size = Pt(9)
    for line in doc.title_lines:
        p = word.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(12)
    widths = (
        [Cm(2.2), Cm(4.0), Cm(0.8), Cm(0.8), Cm(0.8), Cm(3.8), Cm(2.2), Cm(3.4)]
        if doc.with_times
        else [Cm(2.6), Cm(5.6), Cm(0.9), Cm(0.9), Cm(0.9), Cm(4.8), Cm(2.4)]
    )
    for name, blocks in _sections(doc):
        if name:
            p = word.add_paragraph()
            run = p.add_run(name)
            run.bold = True
            run.font.size = Pt(11)
        for block in blocks:
            word.add_paragraph().add_run(block.heading).bold = True
            table = word.add_table(rows=len(block.rows) + 2, cols=len(doc.headers))
            table.style = "Table Grid"
            all_rows = [doc.headers, *block.rows, _total_row(doc, block)]
            for r, values in enumerate(all_rows):
                cells = table.rows[r].cells
                for c, value in enumerate(values):
                    cells[c].text = value
                    cells[c].width = widths[c]
                    if r in (0, len(all_rows) - 1):
                        for run in cells[c].paragraphs[0].runs:
                            run.bold = True
            for c, w in enumerate(widths):
                table.columns[c].width = w
    if not doc.own and not doc.servicing:
        word.add_paragraph("No units entered yet.")
    word.add_paragraph()
    word.add_paragraph(f"Prepared by: {doc.prepared_by}, COD")
    out = BytesIO()
    word.save(out)
    return out.getvalue()


# -- Excel ------------------------------------------------------------------------------


def build_xlsx(doc: Layout) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Allocation"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="E2E8F0")
    for i, w in enumerate([14, 44, 6, 6, 6, 34, 16, 34][: len(doc.headers)], start=1):
        sheet.column_dimensions[chr(64 + i)].width = w
    r = 1
    for line in doc.title_lines:
        sheet.cell(r, 1, line).font = Font(bold=True, size=12)
        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(doc.headers))
        sheet.cell(r, 1).alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    for name, blocks in _sections(doc):
        if name:
            sheet.cell(r, 1, name).font = Font(bold=True, size=11)
            r += 2
        for block in blocks:
            sheet.cell(r, 1, block.heading).font = Font(bold=True)
            r += 1
            for values, bold in [
                (doc.headers, True),
                *[(row, False) for row in block.rows],
                (_total_row(doc, block), True),
            ]:
                for c, value in enumerate(values, start=1):
                    number = c in (3, 4, 5) and value not in ("", "L", "P", "CF")
                    cell = sheet.cell(r, c, float(value) if number else value)
                    cell.border = border
                    cell.font = Font(bold=bold)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if values is doc.headers:
                        cell.fill = fill
                r += 1
            r += 1
    sheet.cell(r + 1, 1, f"Prepared by: {doc.prepared_by}, COD")
    out = BytesIO()
    book.save(out)
    return out.getvalue()


BUILDERS = {"pdf": build_pdf, "docx": build_docx, "xlsx": build_xlsx}
MEDIA = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
